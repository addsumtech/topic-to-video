"""Generate tiny deterministic fixtures: a 2-slide pptx, a 2-page pdf, two PNGs."""
import os, fitz
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))

def make_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    blank = prs.slide_layouts[6]
    for i, (title, note) in enumerate([
        ("Setup the scanner", "Walk the viewer through powering on and homing the scanner."),
        ("Run a scan", "Explain selecting the protocol and pressing start."),
    ], 1):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(1.5))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        s.notes_slide.notes_text_frame.text = note
    out = os.path.join(HERE, "deck.pptx")
    prs.save(out); return out

def make_pdf():
    doc = fitz.open()
    for t in ("Setup the scanner", "Run a scan"):
        page = doc.new_page(width=1280, height=720)
        page.insert_text((100, 360), t, fontsize=48)
    out = os.path.join(HERE, "deck.pdf")
    doc.save(out); return out

def make_images():
    d = os.path.join(HERE, "images"); os.makedirs(d, exist_ok=True)
    doc = fitz.open()
    for i, t in enumerate(("Setup the scanner", "Run a scan"), 1):
        page = doc.new_page(width=1280, height=720)
        page.insert_text((100, 360), t, fontsize=48)
        pix = page.get_pixmap()
        pix.save(os.path.join(d, f"{i:02d}.png"))
    doc.close(); return d

if __name__ == "__main__":
    print(make_pptx()); print(make_pdf()); print(make_images())
