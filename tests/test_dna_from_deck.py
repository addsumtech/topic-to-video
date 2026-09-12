"""dna_from_deck + pptx_theme — the film's Design DNA derived from the deck's OWN theme and frames.

Synthetic decks are built with python-pptx and their theme part PATCHED in the zip, so every branch (Office
default vs custom scheme, serif / mono / humanist faces, sysClr) is exercised without a fixture file. Frames
are PIL-drawn. The two real decks the module was verified on are used when present and skipped otherwise.
"""
import json
import os
import re
import shutil
import sys
import zipfile

import pytest

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "scripts"))

import dna_from_deck as dd            # noqa: E402
import pptx_theme                     # noqa: E402
import theme as th                    # noqa: E402

pptx = pytest.importorskip("pptx")
PIL = pytest.importorskip("PIL")


# ── builders ────────────────────────────────────────────────────────────────────────────────────────
def _make_deck(path, accent1=None, scheme=None, major=None, minor=None, accents=None):
    """A one-slide deck from python-pptx's default template (Office theme), then the theme XML patched."""
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    prs.save(path)
    if not any((accent1, scheme, major, minor, accents)):
        return path
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                x = data.decode("utf-8")
                if scheme:
                    x = re.sub(r'(<a:clrScheme[^>]*name=")[^"]*(")', r"\g<1>%s\2" % scheme, x, count=1)
                for key, val in dict(accents or {}, **({"accent1": accent1} if accent1 else {})).items():
                    x = re.sub(r"<a:%s>.*?</a:%s>" % (key, key),
                               '<a:%s><a:srgbClr val="%s"/></a:%s>' % (key, val.lstrip("#").upper(), key), x, count=1, flags=re.S)
                if major:
                    x = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")', r"\g<1>%s\2" % major, x, count=1)
                if minor:
                    x = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")', r"\g<1>%s\2" % minor, x, count=1)
                data = x.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)
    return path


def _frames(out_dir, ground, accent, n=3, accent_share=0.08, second=None):
    """n solid frames of `ground` with an `accent` band (and an optional smaller `second` band)."""
    from PIL import Image, ImageDraw
    os.makedirs(os.path.join(out_dir, "frames"), exist_ok=True)
    for i in range(1, n + 1):
        im = Image.new("RGB", (320, 180), ground)
        d = ImageDraw.Draw(im)
        d.rectangle([0, 0, 320, int(180 * accent_share)], fill=accent)
        if second:
            d.rectangle([0, 170, 60, 180], fill=second)
        im.save(os.path.join(out_dir, "frames", "%02d.png" % i))


def _board(out_dir, **meta):
    m = {"renderer": "remotion", "title": "t"}
    m.update(meta)
    b = {"meta": m, "slides": [{"index": 0, "shot": "ChatThread", "shot_props": {}, "narration": "one"}]}
    json.dump(b, open(os.path.join(out_dir, "storyboard.json"), "w", encoding="utf-8"))
    return b


def _c(hexstr):
    return th._hex2rgb(hexstr)


# ── pptx_theme ──────────────────────────────────────────────────────────────────────────────────────
def test_the_default_office_theme_is_named_as_default(tmp_path):
    t = pptx_theme.read(_make_deck(str(tmp_path / "d.pptx")))
    assert t["default_theme"], t
    assert t["colors"]["dk1"] == "#000000" and t["colors"]["lt1"] == "#ffffff", "sysClr resolved through lastClr"
    assert t["fonts"]["major"]["latin"] and t["fonts"]["minor"]["latin"]
    assert t["masters"] >= 1 and t["slides"] == 1


def test_a_custom_scheme_is_read_verbatim(tmp_path):
    t = pptx_theme.read(_make_deck(str(tmp_path / "d.pptx"), accent1="#0077cc", scheme="Custom 83",
                                   major="Lora", minor="Roboto", accents={"accent4": "#de7c00"}))
    assert not t["default_theme"]
    assert t["scheme"] == "Custom 83" and t["colors"]["accent1"] == "#0077cc" and t["colors"]["accent4"] == "#de7c00"
    assert t["fonts"]["major"]["latin"] == "Lora" and t["fonts"]["minor"]["latin"] == "Roboto"


def test_master_ground_resolves_through_the_clrmap():
    cols = {"lt1": "#ffffff", "dk1": "#000000", "lt2": "#eeece1", "dk2": "#1f497d"}
    master = '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2"/><p:bg><p:bgRef idx="1001"><a:schemeClr val="bg2"/></p:bgRef></p:bg>'
    assert pptx_theme._master_ground(master, cols) == "#eeece1", "bg2 is lt2 through the map, not 'bg2'"
    assert pptx_theme._master_ground('<p:bg><p:bgPr><a:gradFill/></p:bgPr></p:bg>', cols) is None, "a gradient is not a colour"
    assert pptx_theme._master_ground('<p:cSld/>', cols) == "#ffffff", "no <p:bg>: bg1 shows through"


# ── classification + font map ───────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("name,cls", [("Calibri", "humanist"), ("Calibri Light", "humanist"), ("Helvetica Neue", "humanist"),
                                      ("Lora", "serif"), ("Georgia", "serif"), ("Times New Roman", "serif"),
                                      ("Bodoni 72", "didone"), ("Consolas", "mono"), ("JetBrains Mono", "mono"),
                                      ("Futura", "geometric"), ("Montserrat", "geometric"), ("Comic Sans MS", "rounded"),
                                      ("IBM Plex Sans", "plex"), ("", "humanist"), ("Some Unknown Face", "humanist")])
def test_classify_font(name, cls):
    assert dd.classify_font(name) == cls


def test_every_mapped_face_is_one_the_renderer_bundles():
    faces = dd.loaded_families()
    assert len(faces) >= 10, faces
    for cls, m in dd._CLASS_FACES.items():
        assert m["display"] in faces and m["body"] in faces, (cls, m)


# ── derivation ──────────────────────────────────────────────────────────────────────────────────────
def test_default_theme_light_deck_takes_its_palette_from_the_frames(tmp_path):
    out = str(tmp_path)
    _frames(out, (250, 250, 250), (118, 112, 229), second=(212, 162, 81))
    deck = _make_deck(str(tmp_path / "d.pptx"))
    dna = dd.derive(out, pptx=deck)
    p = dna["provenance"]
    assert p["default_theme"] and p["mode"] == "light" and p["ground"] == "frames"
    assert "frames" in p["accent"], p
    assert dd._hue_distance(_c(dna["accent"]), (118, 112, 229)) < 0.04, "the accent is the frames' dominant saturated hue"
    assert "second hue" in p["accent2"] and dd._hue_distance(_c(dna["accent2"]), (212, 162, 81)) < 0.06
    assert th._contrast(_c(dna["fg"]), _c(dna["bg"])) >= 7.0
    assert th._contrast(_c(dna["accent"]), _c(dna["bg"])) >= 3.0
    assert th._contrast(_c(dna["accent2"]), _c(dna["bg"])) >= 3.0
    assert '"Inter"' in dna["fontDisplay"] and "Noto Sans CJK SC" in dna["fontBody"] and dna["fontBody"].strip().endswith("sans-serif")


def test_custom_theme_keeps_its_accent_and_classifies_its_faces(tmp_path):
    out = str(tmp_path)
    _frames(out, (10, 20, 60), (200, 200, 200), accent_share=0.02)                 # dark deck, no strong hue
    # every accent set explicitly: the template's own accents 4-6 would otherwise compete for "most distant"
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#e2603f", scheme="Brand", major="Lora", minor="Roboto",
                      accents={"accent2": "#cfd2d9", "accent3": "#196b24", "accent4": "#f79646",
                               "accent5": "#c0504d", "accent6": "#e8a33d"})
    dna = dd.derive(out, pptx=deck)
    p = dna["provenance"]
    assert not p["default_theme"] and p["mode"] == "dark" and p["accent"] == "theme accent1"
    assert dd._hue_distance(_c(dna["accent"]), _c("#e2603f")) < 0.03, "hue kept while nudged to 3:1"
    assert "most distant" in p["accent2"] and dd._hue_distance(_c(dna["accent2"]), _c("#196b24")) < 0.04, \
        "the grey accent2 is skipped; the green is the most distant hue"
    assert '"Playfair Display"' in dna["fontDisplay"] and '"Source Serif 4"' not in dna["fontBody"] and '"Inter"' in dna["fontBody"]
    assert dna["atoms"]["label"] == "kicker" and dna["uppercaseLabels"] is True and dna["radius"] == 2


def test_a_mono_theme_speaks_the_terminal_grammar(tmp_path):
    out = str(tmp_path)
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#3dd6a4", scheme="Term", major="Consolas", minor="Consolas")
    dna = dd.derive(out, pptx=deck, mode="dark")                                    # no frames: theme only
    assert dna["uniformType"] is True and dna["motion"] == "type" and dna["atoms"]["list"] == "tree"
    assert '"JetBrains Mono"' in dna["fontDisplay"] and '"JetBrains Mono"' in dna["fontBody"]
    assert dna["provenance"]["source"].endswith("d.pptx") and dna["provenance"]["ground"] != "frames"


def test_theme_only_dark_takes_dk2_as_the_ground(tmp_path):
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#6dcbf4", scheme="Codex", accents={"dk2": "#00264f"})
    dna = dd.derive(str(tmp_path), pptx=deck, mode="dark")
    assert dna["provenance"]["ground"] == "theme dk2" and dna["bg"].startswith("#00")
    assert th._contrast(_c(dna["fg"]), _c(dna["bg"])) >= 7.0


def test_frames_only_works_and_nothing_at_all_fails_loudly(tmp_path):
    out = str(tmp_path / "a"); os.makedirs(out)
    _frames(out, (0, 83, 143), (118, 178, 215))
    dna = dd.derive(out, pptx=None)
    assert dna["provenance"]["source"] == "frames only" and dna["provenance"]["mode"] == "dark"
    with pytest.raises(SystemExit):
        dd.derive(str(tmp_path / "b"), pptx=None)


def test_derivation_is_deterministic(tmp_path):
    out = str(tmp_path)
    _frames(out, (250, 250, 250), (118, 112, 229))
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#0077cc", scheme="X")
    assert dd.derive(out, pptx=deck) == dd.derive(out, pptx=deck)


def test_on_accent_is_picked_by_true_contrast(tmp_path):
    out = str(tmp_path)
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#f7cb59", scheme="Y")   # a yellow accent wants dark ink
    dna = dd.derive(out, pptx=deck, mode="light")
    acc, on = _c(dna["accent"]), _c(dna["onAccent"])
    assert th._contrast(on, acc) >= th._contrast((255, 255, 255) if on != (255, 255, 255) else (17, 17, 17), acc)


# ── the CLI: what it writes, what it refuses ────────────────────────────────────────────────────────
def test_main_writes_the_dna_and_records_the_source(tmp_path, capsys):
    out = str(tmp_path)
    _frames(out, (250, 250, 250), (118, 112, 229))
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#0077cc", scheme="X")
    _board(out, design_dna="product")                                                # a preset NAME is replaced
    assert dd.main([out, "--pptx", deck]) == 0
    b = json.load(open(os.path.join(out, "storyboard.json"), encoding="utf-8"))
    assert isinstance(b["meta"]["design_dna"], dict) and b["meta"]["design_dna"]["id"] == "deck"
    assert b["meta"]["design_dna_source"] == "deck" and b["meta"]["source"] == os.path.abspath(deck)
    assert "wrote meta.design_dna" in capsys.readouterr().out


def test_main_refuses_to_overwrite_an_authored_object_without_force(tmp_path):
    out = str(tmp_path)
    _frames(out, (250, 250, 250), (118, 112, 229))
    _board(out, design_dna={"id": "terminal", "bg": "#000"})
    assert dd.main([out]) == 1
    b = json.load(open(os.path.join(out, "storyboard.json"), encoding="utf-8"))
    assert b["meta"]["design_dna"] == {"id": "terminal", "bg": "#000"}, "left exactly as authored"
    assert dd.main([out, "--force"]) == 0
    assert json.load(open(os.path.join(out, "storyboard.json"), encoding="utf-8"))["meta"]["design_dna"]["id"] == "deck"


def test_dry_run_writes_nothing_and_finds_the_deck_through_meta_source(tmp_path, capsys):
    out = str(tmp_path)
    _frames(out, (250, 250, 250), (118, 112, 229))
    deck = _make_deck(str(tmp_path / "d.pptx"), accent1="#0077cc", scheme="X", major="Lora")
    before = _board(out, source=deck)
    assert dd.main([out, "--dry-run"]) == 0
    assert json.load(open(os.path.join(out, "storyboard.json"), encoding="utf-8")) == before
    assert "Lora (serif) -> Playfair Display" in capsys.readouterr().out, "the deck was found via meta.source"


# ── validate: the nudge ─────────────────────────────────────────────────────────────────────────────
def test_validate_nudges_a_custom_themed_pptx_that_wears_a_preset(tmp_path):
    import validate
    out = str(tmp_path)
    custom = _make_deck(str(tmp_path / "custom.pptx"), accent1="#0077cc", scheme="Custom 83", major="Calibri Light")
    default = _make_deck(str(tmp_path / "office.pptx"))
    hits = lambda b: [w for w in validate.evaluate(b, out, pre_tts=True)["warnings"] if "OWN theme" in w]   # noqa: E731
    assert hits(_board(out, source=custom, design_dna="product")), "custom theme + preset: nudged"
    assert hits(_board(out, source=custom)), "custom theme + no DNA: nudged"
    assert not hits(_board(out, source=default, design_dna="product")), "Office default: nothing to carry"
    assert not hits(_board(out, source=custom, design_dna={"id": "deck", "bg": "#00538f"})), "an object is a decision"


def test_ingest_records_the_source_deck(tmp_path):
    import ingest
    deck = _make_deck(str(tmp_path / "d.pptx"))
    if not shutil.which("soffice"):
        pytest.skip("soffice")
    out = str(tmp_path / "out")
    data = ingest.ingest(deck, out)
    assert data["meta"]["source"] == os.path.abspath(deck)


# ── the two real decks, when present ────────────────────────────────────────────────────────────────
_REAL = {
    "leiden": "/Users/donghanglyu/Downloads/4D_cine_documents/2026-04-13 4d FBWH explanation for Leiden.pptx",
    "workbuddy": "/Users/donghanglyu/WorkBuddy/2026-07-24-08-34-58/workbuddy-intro.pptx",
}


def _readable(path):
    """PRESENT IS NOT THE SAME AS READABLE on this platform, and the difference is a red suite.

    macOS TCC can withdraw access to ~/Downloads and ~/Desktop part-way through a session: the path still
    EXISTS, so an `os.path.exists` guard lets the test through, and python-pptx then fails with
    "Package not found at ..." — which reads like a missing deck and is a permission. Open it.
    """
    try:
        with open(path, "rb") as fh:
            fh.read(4)
        return True
    except OSError:
        return False


def test_real_leiden_deck_is_a_custom_blue_theme():
    if not _readable(_REAL["leiden"]):
        pytest.skip("real deck not readable on this machine")
    t = pptx_theme.read(_REAL["leiden"])
    assert not t["default_theme"] and t["colors"]["accent1"] == "#0077cc" and t["fonts"]["major"]["latin"] == "Calibri Light"


def test_real_workbuddy_deck_is_office_default():
    if not _readable(_REAL["workbuddy"]):
        pytest.skip("real deck not readable on this machine")
    t = pptx_theme.read(_REAL["workbuddy"])
    assert t["default_theme"] and t["fonts"]["major"]["latin"] == "Calibri"


def test_every_derived_font_stack_names_a_bundled_cjk_face_before_a_system_one():
    """The bundling was done for determinism: the same film renders the same on any machine. A DNA derived
    from a deck builds its OWN stacks and did not name the bundled subsets at all, so the one path a real
    Chinese film takes fell through to whatever the render machine had installed — PingFang on a Mac,
    something else everywhere else. `runtime/dna.tsx`'s presets had been fixed; this file had not."""
    import re
    import sys
    import os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
    import dna_from_deck as d

    bundled = d.loaded_families()
    assert {"Noto Sans SC", "Noto Serif SC"} <= bundled, (
        "fonts.ts no longer bundles the CJK subsets — %s" % sorted(bundled))
    for face, generic in (("Playfair Display", "Georgia, serif"), ("Inter", "system-ui, sans-serif"),
                          ("JetBrains Mono", "ui-monospace, monospace"),
                          ("Source Serif 4", "Georgia, serif")):
        stack = d._stack(face, generic)
        names = re.findall(r'"([^"]+)"', stack)
        cjk = [n for n in names if "SC" in n or "YaHei" in n or "Hiragino" in n or "SimSun" in n]
        assert cjk, stack
        assert cjk[0] in bundled, (
            "%s falls through to a SYSTEM face (%s) before any bundled one: %s" % (face, cjk[0], stack))


def test_the_serif_worlds_take_the_serif_subset():
    """A serif Latin face over a sans Chinese one is two typefaces pretending to be one."""
    import sys
    import os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
    import dna_from_deck as d
    assert '"Noto Serif SC"' in d._stack("Playfair Display", "Georgia, serif")
    assert '"Noto Sans SC"' in d._stack("Inter", "system-ui, sans-serif")


def test_validate_warns_when_a_cjk_film_names_no_bundled_face():
    """The gate for the same hole one level up: a hand-authored or older DNA can still list only system
    families, and then the film's typography is whatever the render machine happens to have. Controlled in
    both directions, and silent on an English film — the faces are irrelevant there."""
    import sys
    import os
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
    import validate as v

    def board(stack, lang="zh", copy="交付可验收的成果"):
        return {"meta": {"renderer": "remotion", "title": "t", "language": lang,
                         "design_dna": {"id": "x", "fontDisplay": stack, "fontBody": stack,
                                        "fontMono": stack, "bg": "#fff", "fg": "#111"}},
                "slides": [{"index": 0, "shot": "ChatThread", "shot_props": {"title": copy},
                            "narration": copy}]}

    def warns(b):
        return [w for w in v.evaluate(b, "/tmp", pre_tts=True)["warnings"] if "BUNDLED CJK face" in w]

    sysonly = '"Inter", "PingFang SC", "Microsoft YaHei", system-ui, sans-serif'
    bundled = '"Inter", "Noto Sans SC", "PingFang SC", system-ui, sans-serif'
    assert warns(board(sysonly)), "a CJK film with only system faces must be told"
    assert not warns(board(bundled)), "naming the bundled face must silence it"
    assert not warns(board(sysonly, "en", "An accepted result")), "an English film must never see this"
