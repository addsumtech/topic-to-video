"""Turn any supported input into frames/*.png + a skeleton storyboard.json."""
import argparse, os, glob, shutil, tempfile
import fitz  # PyMuPDF
import _common
import storyboard as sb

IMAGE_EXTS = (".png", ".jpg", ".jpeg")

def _render_pdf(pdf_path, frames_dir, dpi=150):
    doc = fitz.open(pdf_path)
    paths = []
    for i, page in enumerate(doc, 1):
        pix = page.get_pixmap(dpi=dpi)
        out = os.path.join(frames_dir, f"{i:02d}.png")
        pix.save(out)
        paths.append(out)
    doc.close()
    return paths

def _pptx_to_pdf(pptx_path, workdir):
    soffice = _common.find_soffice()
    if soffice is None:
        raise RuntimeError("LibreOffice (soffice) not found; cannot convert .pptx")
    _common.run([soffice, "--headless", "--convert-to", "pdf",
                 "--outdir", workdir, pptx_path])
    stem = os.path.splitext(os.path.basename(pptx_path))[0]
    return os.path.join(workdir, stem + ".pdf")

def _pptx_notes(pptx_path):
    from pptx import Presentation
    notes = []
    for slide in Presentation(pptx_path).slides:
        txt = ""
        if slide.has_notes_slide:
            txt = slide.notes_slide.notes_text_frame.text or ""
        notes.append(txt.strip())
    return notes

def ingest(input_path, out_dir, dpi=150):
    frames_dir = os.path.join(out_dir, "frames")
    os.makedirs(frames_dir, exist_ok=True)
    notes = None

    if os.path.isdir(input_path):
        imgs = sorted(f for f in glob.glob(os.path.join(input_path, "*"))
                      if f.lower().endswith(IMAGE_EXTS))
        frame_paths = []
        for i, src in enumerate(imgs, 1):
            dst = os.path.join(frames_dir, f"{i:02d}.png")
            shutil.copyfile(src, dst); frame_paths.append(dst)
    elif input_path.lower().endswith(".pdf"):
        frame_paths = _render_pdf(input_path, frames_dir, dpi)
    elif input_path.lower().endswith(".pptx"):
        with tempfile.TemporaryDirectory() as wd:
            pdf = _pptx_to_pdf(input_path, wd)
            frame_paths = _render_pdf(pdf, frames_dir, dpi)
        notes = _pptx_notes(input_path)
    else:
        raise ValueError(f"unsupported input: {input_path}")

    if not frame_paths:
        raise ValueError(f"no slides found in {input_path}")

    data = sb.new_storyboard()
    for i, fp in enumerate(frame_paths, 1):
        note = notes[i - 1] if notes and i - 1 < len(notes) else ""
        rel = os.path.relpath(fp, out_dir)
        data["slides"].append(sb.new_slide(i, rel, narration="", notes=note))
    sb.save(data, os.path.join(out_dir, "storyboard.json"))
    return data

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("out_dir")
    ap.add_argument("--dpi", type=int, default=150)
    a = ap.parse_args()
    data = ingest(a.input, a.out_dir, a.dpi)
    print(f"Ingested {len(data['slides'])} slides -> {a.out_dir}/storyboard.json")
    if any(s["notes"] for s in data["slides"]):
        print("Speaker notes found and stored in each slide's 'notes' field.")

if __name__ == "__main__":
    main()
